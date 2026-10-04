from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.shared import Pt

SECTIONS = (
    ("summary", "Professional Summary"),
    ("skills", "Skills"),
    ("experience", "Experience"),
    ("education", "Education"),
    ("certifications", "Certifications"),
    ("projects", "Projects"),
)


def draft_to_text(draft: dict[str, str]) -> str:
    lines = [draft.get("name", "").strip()]
    contact = [draft.get(key, "").strip() for key in ("email", "phone", "linkedin")]
    lines.extend(value for value in contact if value)
    for key, heading in SECTIONS:
        value = draft.get(key, "").strip()
        if value:
            lines.extend(["", heading, value])
    return "\n".join(lines)


def draft_to_docx(draft: dict[str, str]) -> bytes:
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)

    name = draft.get("name", "").strip()
    if name:
        document.add_heading(name, level=0)
    contact = " | ".join(
        draft.get(key, "").strip()
        for key in ("email", "phone", "linkedin")
        if draft.get(key, "").strip()
    )
    if contact:
        document.add_paragraph(contact)

    for key, heading in SECTIONS:
        value = draft.get(key, "").strip()
        if not value:
            continue
        document.add_heading(heading, level=1)
        for line in value.splitlines():
            clean = line.strip()
            if not clean:
                continue
            if clean.startswith(("-", "•", "*")):
                document.add_paragraph(clean.lstrip("-•* ").strip(), style="List Bullet")
            else:
                document.add_paragraph(clean)

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()