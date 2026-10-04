from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Any

import pdfplumber
import spacy
from docx import Document

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
NLP = spacy.blank("en")

HEADING_ALIASES = {
    "summary": "summary",
    "professional summary": "summary",
    "profile": "summary",
    "objective": "summary",
    "experience": "experience",
    "professional experience": "experience",
    "work experience": "experience",
    "employment history": "experience",
    "work history": "experience",
    "skills": "skills",
    "technical skills": "skills",
    "core skills": "skills",
    "education": "education",
    "certifications": "certifications",
    "certificates": "certifications",
    "projects": "projects",
    "selected projects": "projects",
}

EMAIL_RE = re.compile(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}")
PHONE_RE = re.compile(r"(?<!\w)(?:\+?\d[\d ().-]{7,}\d)(?!\w)")
LINKEDIN_RE = re.compile(r"(?:https?://)?(?:www\.)?linkedin\.com/in/[A-Za-z0-9._%-]+", re.I)


def tokenize(text: str) -> list[str]:
    """Use spaCy's English tokenizer without requiring a downloaded model."""
    return [
        token.text.lower()
        for token in NLP(text)
        if not token.is_space and not token.is_punct and not token.is_stop
    ]


def split_sections(text: str) -> dict[str, str]:
    sections: dict[str, list[str]] = {}
    current: str | None = None
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            if current:
                sections.setdefault(current, []).append("")
            continue
        normalized = re.sub(r"[:\s]+$", "", line.lower())
        heading = HEADING_ALIASES.get(normalized)
        if heading:
            current = heading
            sections.setdefault(current, [])
            continue
        if current:
            sections.setdefault(current, []).append(line)
    return {key: "\n".join(lines).strip() for key, lines in sections.items()}


def draft_from_text(text: str) -> dict[str, Any]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    contact_line = " ".join(lines[:8])
    name = next(
        (
            line for line in lines[:5]
            if not EMAIL_RE.search(line)
            and not PHONE_RE.search(line)
            and not LINKEDIN_RE.search(line)
            and len(line) < 80
        ),
        "",
    )
    sections = split_sections(text)
    return {
        "name": name,
        "email": (EMAIL_RE.search(contact_line).group(0) if EMAIL_RE.search(contact_line) else ""),
        "phone": (PHONE_RE.search(contact_line).group(0).strip() if PHONE_RE.search(contact_line) else ""),
        "linkedin": (LINKEDIN_RE.search(contact_line).group(0) if LINKEDIN_RE.search(contact_line) else ""),
        "summary": sections.get("summary", ""),
        "skills": sections.get("skills", ""),
        "experience": sections.get("experience", ""),
        "education": sections.get("education", ""),
        "certifications": sections.get("certifications", ""),
        "projects": sections.get("projects", ""),
    }


def _extract_docx(data: bytes) -> tuple[str, list[str]]:
    document = Document(io.BytesIO(data))
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    formatting_issues: list[str] = []
    font_names = {
        run.font.name.strip()
        for paragraph in document.paragraphs
        for run in paragraph.runs
        if run.font.name and run.font.name.strip()
    }
    accepted_fonts = {
        "arial", "calibri", "aptos", "times new roman", "cambria",
        "georgia", "helvetica", "verdana", "tahoma",
    }
    unusual = sorted(font for font in font_names if font.lower() not in accepted_fonts)
    if unusual:
        formatting_issues.append(f"Non-standard font detected: {', '.join(unusual[:3])}.")
    if document.tables:
        formatting_issues.append("The DOCX contains tables; check that important text reads in the intended order.")
    return "\n".join(parts), formatting_issues


def _extract_pdf(data: bytes) -> tuple[str, list[str]]:
    pages: list[str] = []
    formatting_issues: list[str] = []
    possible_columns = 0
    has_images = False
    with pdfplumber.open(io.BytesIO(data)) as document:
        for page in document.pages:
            pages.append(page.extract_text(layout=True) or "")
            has_images = has_images or bool(page.images)
            words = page.extract_words() or []
            rows: dict[int, list[float]] = {}
            for word in words:
                row = round(float(word.get("top", 0)) / 3)
                rows.setdefault(row, []).append(float(word.get("x0", 0)))
            for positions in rows.values():
                ordered = sorted(positions)
                if any(right - left > 115 for left, right in zip(ordered, ordered[1:])):
                    possible_columns += 1
    if possible_columns >= 6:
        formatting_issues.append("Text appears to use multiple columns; ATS readers may mix the reading order.")
    if has_images:
        formatting_issues.append("The PDF contains image content; make sure important details are also selectable text.")
    return "\n".join(pages), formatting_issues


def extract_resume(filename: str, data: bytes) -> dict[str, Any]:
    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError("This file is over 10 MB. Choose a smaller resume file.")

    extension = Path(filename).suffix.lower()
    if extension == ".txt":
        text = data.decode("utf-8-sig", errors="replace")
        formatting_issues: list[str] = []
    elif extension == ".pdf":
        text, formatting_issues = _extract_pdf(data)
    elif extension == ".docx":
        text, formatting_issues = _extract_docx(data)
    else:
        raise ValueError("Choose a PDF, DOCX, or TXT resume.")

    if len(text.strip()) < 40:
        raise ValueError("We could not find enough readable text. Try a text-based PDF or DOCX.")
    return {
        "name": filename,
        "raw_text": text.strip(),
        "formatting_issues": formatting_issues,
        "draft": draft_from_text(text),
    }